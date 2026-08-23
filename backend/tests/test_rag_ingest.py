"""RAG multi-format document ingestion + topic chat reference injection."""

import io

import pytest


@pytest.fixture()
def offline_rag(monkeypatch):
    from backend.app.services.rag import rag_service

    class FakeStore:
        def __init__(self):
            self.docs = []

        def add_documents(self, docs):
            self.docs.extend(docs)

        def similarity_search(self, query, k=4):
            tokens = [token for token in query.split() if token]
            hits = [doc for doc in self.docs if any(token in doc.page_content for token in tokens)]
            return (hits or list(self.docs))[:k]

    store = FakeStore()
    monkeypatch.setattr(rag_service, "_check_init", lambda: None)
    monkeypatch.setattr(rag_service, "vector_store", store)
    monkeypatch.setattr(rag_service, "_add_accounted", lambda splits: store.add_documents(splits) or len(splits))
    return rag_service


def _ingest(client, filename: str, payload: bytes):
    return client.post("/api/rag/ingest", files={"file": (filename, payload, "application/octet-stream")})


def _search_all(service, query: str):
    import asyncio

    return asyncio.run(service.search(query, 5))


def _xlsx_payload() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "月度产量"
    sheet.append(["月份", "产量"])
    sheet.append(["1月", 120])
    extra = book.create_sheet("质量问题")
    extra.append(["缺陷类型", "次数"])
    extra.append(["密封渗漏", 8])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _docx_payload() -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph("泵站机组故障背景：近三个月故障频发。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "月均故障"
    table.cell(1, 1).text = "6 次"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _pptx_payload() -> bytes:
    from pptx import Presentation
    from pptx.util import Cm

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Cm(2), Cm(2), Cm(12), Cm(4))
    box.text_frame.text = "初步对策：统一更换机械密封件"
    slide.notes_slide.notes_text_frame.text = "备注：车间已同意试点两台机组"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _pdf_payload() -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = b"BT /F1 24 Tf 72 720 Td (QC-TOPIC-REFERENCE-PDF) Tj ET"
    objects[3] = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream)
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % index + body + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, xref_at)
    return bytes(out)


def test_ingest_markdown(client, offline_rag):
    response = _ingest(client, "选题思路.md", "# 课题方向\n降低密封件泄漏率".encode())
    assert response.status_code == 200 and response.json()["chunks"] >= 1


def test_ingest_txt_with_gb18030_encoding(client, offline_rag):
    response = _ingest(client, "值班记录.txt", "泵站夜班巡检发现渗漏。".encode("gb18030"))
    assert response.status_code == 200


def test_ingest_csv_and_xlsx(client, offline_rag):
    csv_response = _ingest(client, "缺陷统计.csv", "月份,缺陷数\n1月,3\n2月,5\n".encode())
    assert csv_response.status_code == 200 and csv_response.json()["chunks"] >= 1

    xlsx_response = _ingest(client, "生产台账.xlsx", _xlsx_payload())
    assert xlsx_response.status_code == 200
    # 两个非空工作表各自成块，且内容确实入库
    assert xlsx_response.json()["chunks"] >= 2
    found = _search_all(offline_rag, "密封渗漏")
    assert any("密封渗漏" in doc.page_content for doc in found)


def test_ingest_docx_pptx_pdf(client, offline_rag):
    docx_response = _ingest(client, "课题报告.docx", _docx_payload())
    assert docx_response.status_code == 200 and docx_response.json()["chunks"] >= 1

    pptx_response = _ingest(client, "汇报.pptx", _pptx_payload())
    assert pptx_response.status_code == 200 and pptx_response.json()["chunks"] >= 1

    pdf_response = _ingest(client, "规程.pdf", _pdf_payload())
    assert pdf_response.status_code == 200 and pdf_response.json()["chunks"] >= 1
    texts = [doc.page_content for doc in _search_all(offline_rag, "REFERENCE")]
    assert any("QC-TOPIC-REFERENCE-PDF" in text for text in texts)


def test_ingest_rejects_unsupported_and_empty(client, offline_rag):
    zip_response = _ingest(client, "打包.zip", b"PK\x03\x04")
    assert zip_response.status_code == 400 and zip_response.json()["code"] == "DATA_FORMAT_UNSUPPORTED"

    legacy = _ingest(client, "老文档.doc", b"\xd0\xcf\x11\xe0")
    assert legacy.status_code == 400 and legacy.json()["code"] == "DATA_FORMAT_UNSUPPORTED"

    empty = _ingest(client, "空白.txt", "   \n\t\n".encode())
    assert empty.status_code == 400 and empty.json()["code"] == "DOCUMENT_TEXT_EMPTY"


def test_topic_chat_injects_references(client, monkeypatch):
    from backend.app.api.endpoints import topic as topic_endpoint

    captured: dict = {}

    async def fake_chat(messages, role=None):
        captured["messages"] = messages
        return "建议聚焦密封件泄漏率"

    class Doc:
        page_content = "去年课题：泵站故障率下降30%\n详见台账"

    async def fake_search(_query, _k):
        return [Doc(), Doc()]

    monkeypatch.setattr(topic_endpoint, "chat", fake_chat)
    monkeypatch.setattr(topic_endpoint.rag_service, "search", fake_search)
    result = client.post("/api/topic/chat", json={"message": "泵站经常故障怎么选题"}).json()
    assert result["response"].startswith("建议") and result["references_used"] == 1
    system = captured["messages"][0]["content"]
    assert "参考资料" in system and "泵站故障率下降30%" in system


def test_topic_chat_survives_rag_failure(client, monkeypatch):
    from backend.app.api.endpoints import topic as topic_endpoint

    async def fake_chat(messages, role=None):
        return "请补充对象与目标值"

    async def broken_search(_query, _k):
        raise RuntimeError("embedding down")

    monkeypatch.setattr(topic_endpoint, "chat", fake_chat)
    monkeypatch.setattr(topic_endpoint.rag_service, "search", broken_search)
    result = client.post("/api/topic/chat", json={"message": "怎么选题"}).json()
    assert result["references_used"] == 0 and result["response"] == "请补充对象与目标值"
