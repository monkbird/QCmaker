"""Extract plain-text Documents from common office formats for RAG ingestion."""

from __future__ import annotations

import csv
import io
import logging
from pathlib import Path

from langchain_core.documents import Document

from backend.app.core.errors import AppException

logger = logging.getLogger(__name__)

TEXT_SUFFIXES = {".txt", ".md", ".markdown"}
TABULAR_SUFFIXES = {".csv", ".xlsx", ".xls"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | TABULAR_SUFFIXES | {".pdf", ".docx", ".pptx"}

MAX_TABLE_ROWS = 1000


def decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise AppException("DATA_FORMAT_UNSUPPORTED", "文本编码无法识别，请另存为 UTF-8 后重试", 400)


def _grid_to_text(rows: list[list[str]]) -> str:
    lines: list[str] = []
    for row in rows:
        cells = [str(cell).replace("\n", " ").strip() for cell in row]
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _frame_documents(path: Path, reader) -> list[Document]:
    import pandas as pd

    try:
        sheets = reader()
    except Exception as exc:
        raise AppException("DATA_FORMAT_UNSUPPORTED", f"表格解析失败：{exc}", 400) from exc
    documents: list[Document] = []
    for name, frame in sheets.items():
        frame = frame.astype(object).where(pd.notna(frame), "")
        truncated = len(frame) > MAX_TABLE_ROWS
        if truncated:
            frame = frame.iloc[:MAX_TABLE_ROWS]
        rows = [list(map(str, row)) for row in frame.itertuples(index=False)]
        header = [str(column) for column in frame.columns]
        body = _grid_to_text([header, *rows])
        if not body.strip():
            continue
        note = f"\n（数据过长，仅保留前 {MAX_TABLE_ROWS} 行）" if truncated else ""
        documents.append(Document(page_content=f"【{name}】\n{body}{note}", metadata={"sheet": str(name)}))
    return documents


def _text_documents(path: Path) -> list[Document]:
    return [Document(page_content=decode_text(path.read_bytes()), metadata={})]


def _csv_documents(path: Path) -> list[Document]:
    decoded = decode_text(path.read_bytes())
    rows = list(csv.reader(io.StringIO(decoded)))
    if len(rows) > MAX_TABLE_ROWS:
        rows = rows[:MAX_TABLE_ROWS] + [["…（数据过长，已截断）"]]
    body = _grid_to_text(rows)
    if not body.strip():
        return []
    return [Document(page_content=body, metadata={})]


def _excel_documents(path: Path) -> list[Document]:
    def read():
        import pandas as pd

        return pd.read_excel(path, sheet_name=None)

    return _frame_documents(path, read)


def _docx_documents(path: Path) -> list[Document]:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
    for table in doc.tables:
        grid = [[cell.text for cell in row.cells] for row in table.rows]
        rendered = _grid_to_text(grid)
        if rendered:
            parts.append(rendered)
    if not parts:
        return []
    return [Document(page_content="\n".join(parts), metadata={})]


def _pptx_documents(path: Path) -> list[Document]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    documents: list[Document] = []
    for index, slide in enumerate(presentation.slides, start=1):
        chunks: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                chunks.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                grid = [[cell.text for cell in row.cells] for row in shape.table.rows]
                rendered = _grid_to_text(grid)
                if rendered:
                    chunks.append(rendered)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                chunks.append(notes)
        joined = "\n".join(chunks).strip()
        if joined:
            documents.append(Document(page_content=joined, metadata={"slide": index}))
    return documents


def _pdf_documents(path: Path) -> list[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    documents: list[Document] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            documents.append(Document(page_content=text, metadata={"page": index}))
    return documents


def extract_documents(file_path: str | Path, original_filename: str) -> list[Document]:
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise AppException(
            "DATA_FORMAT_UNSUPPORTED",
            "仅支持 TXT / Markdown / PDF / Word(.docx) / Excel(.xlsx/.xls) / CSV / PPT(.pptx)",
            400,
        )
    loaders = {
        ".pdf": _pdf_documents,
        ".docx": _docx_documents,
        ".pptx": _pptx_documents,
        ".xlsx": _excel_documents,
        ".xls": _excel_documents,
        ".csv": _csv_documents,
    }
    try:
        documents = loaders.get(suffix, _text_documents)(path)
    except AppException:
        raise
    except Exception as exc:
        logger.warning("document parse failed for %s: %s", original_filename, exc)
        raise AppException("DATA_FORMAT_UNSUPPORTED", f"文档解析失败：{exc}", 400) from exc
    documents = [document for document in documents if document.page_content.strip()]
    if not documents:
        raise AppException("DOCUMENT_TEXT_EMPTY", "未能从文档中提取到文本内容", 400)
    for document in documents:
        document.metadata["source"] = original_filename
    return documents
